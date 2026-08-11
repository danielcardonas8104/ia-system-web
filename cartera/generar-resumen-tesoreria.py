from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/ia-system-web/cartera/Polizas-y-sumas-aseguradas-tesoreria-2026-07-28.xlsx"

F = "Arial"
INK = "FF111827"
WHITE = "FFFFFFFF"
HDR_FILL = PatternFill("solid", fgColor="FF334155")
BAND = PatternFill("solid", fgColor="FFF3F4F6")
INPUT_FILL = PatternFill("solid", fgColor="FFFFFF00")
SEC_FILL = PatternFill("solid", fgColor="FFE5E7EB")
BLUE = "FF0000FF"

thin = Side(style="thin", color="FFD1D5DB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TOPRULE = Border(top=Side(style="medium", color="FF334155"))

MXN = '$#,##0;($#,##0);-'
NUM = '#,##0.00;(#,##0.00);-'
RATE = '#,##0.000000'

POLIZAS = [
    ("89793", "DANIEL CARDONA SANCHEZ", "DANIEL CARDONA SANCHEZ", "Titular",
     "BIENESTAR PRUDENTIAL 20 AÑOS SIN INVALIDEZ", "02/07/2024", "EN VIGOR", "UDIS",
     150000, 150000, 150000, 40000),
    ("102239", "DANIEL CARDONA SANCHEZ", "DANIEL CARDONA SANCHEZ", "Titular",
     "PERSONALIZA PRU 20 AÑOS", "30/10/2025", "EN VIGOR", "UDIS",
     60000, 200000, 200000, 0),
    ("102261", "DANIEL CARDONA SANCHEZ", "DANIEL CARDONA SANCHEZ", "Titular",
     "GARANTÍA PRUDENTIAL 20 AÑOS SIN INVALIDEZ", "30/10/2025", "EN VIGOR", "UDIS",
     225000, 225000, 225000, 0),
    ("102264", "DANIEL CARDONA SANCHEZ", "DANIEL CARDONA SANCHEZ", "Titular",
     "BIENESTAR PRUDENTIAL 5 AÑOS SIN INVALIDEZ", "30/10/2025", "EN VIGOR", "DOLARES",
     113300, 113300, 113300, 0),
    ("92834", "MARTHA LETICIA GOSSEREZ ROJAS", "MARTHA LETICIA GOSSEREZ ROJAS", "Grupo familiar",
     "BIENESTAR PRUDENTIAL 10 AÑOS SIN INVALIDEZ", "31/10/2024", "EN VIGOR", "UDIS",
     250000, 250000, 250000, 27500),
    ("100631", "MARTHA LETICIA GOSSEREZ ROJAS", "ALDO CARDONA GOSSEREZ", "Grupo familiar",
     "PERSONALIZA PRU 20 AÑOS", "25/09/2025", "EN VIGOR", "UDIS",
     60000, 0, 370000, 0),
]

COBERTURAS = [
    ("89793", "DANIEL CARDONA SANCHEZ", "BIENESTAR PRUDENTIAL 20 AÑOS SIN INVALIDEZ (UDIS)", "BASICA", 150000, "UDIS", "PAGO DE PRIMAS EN CURSO"),
    ("89793", "DANIEL CARDONA SANCHEZ", "CIRUGÍAS BRONCE SIN INVALIDEZ (UDIS)", "SALUD", 40000, "UDIS", "PAGO DE PRIMAS EN CURSO"),
    ("89793", "DANIEL CARDONA SANCHEZ", "INVALIDEZ TOTAL Y PERMANENTE", "INVALIDEZ", 150000, "UDIS", "no disponible"),
    ("89793", "DANIEL CARDONA SANCHEZ", "MUERTE ACCIDENTAL O PÉRDIDAS ORGÁNICAS", "ACCIDENTAL", 150000, "UDIS", "no disponible"),
    ("102239", "DANIEL CARDONA SANCHEZ", "PERSONALIZA PRU 20 AÑOS (UDIS)", "BASICA", 60000, "UDIS", "PAGO DE PRIMAS EN CURSO"),
    ("102239", "DANIEL CARDONA SANCHEZ", "EXCESO PERSONALIZA PRU 20 AÑOS (UDIS)", "EXCESO", 0, "UDIS", "PAGO DE PRIMAS EN CURSO"),
    ("102239", "DANIEL CARDONA SANCHEZ", "INVALIDEZ TOTAL Y PERMANENTE", "INVALIDEZ", 200000, "UDIS", "no disponible"),
    ("102239", "DANIEL CARDONA SANCHEZ", "MUERTE ACCIDENTAL O PÉRDIDAS ORGÁNICAS", "ACCIDENTAL", 200000, "UDIS", "no disponible"),
    ("102261", "DANIEL CARDONA SANCHEZ", "GARANTÍA PRUDENTIAL 20 AÑOS SIN INVALIDEZ (UDIS)", "BASICA", 225000, "UDIS", "PAGO DE PRIMAS EN CURSO"),
    ("102261", "DANIEL CARDONA SANCHEZ", "INVALIDEZ TOTAL Y PERMANENTE", "INVALIDEZ", 225000, "UDIS", "no disponible"),
    ("102261", "DANIEL CARDONA SANCHEZ", "MUERTE ACCIDENTAL O PÉRDIDAS ORGÁNICAS", "ACCIDENTAL", 225000, "UDIS", "no disponible"),
    ("102264", "DANIEL CARDONA SANCHEZ", "BIENESTAR PRUDENTIAL 5 AÑOS SIN INVALIDEZ (USD)", "BASICA", 113300, "DOLARES", "PAGO DE PRIMAS EN CURSO"),
    ("102264", "DANIEL CARDONA SANCHEZ", "INVALIDEZ TOTAL Y PERMANENTE", "INVALIDEZ", 113300, "DOLARES", "no disponible"),
    ("102264", "DANIEL CARDONA SANCHEZ", "MUERTE ACCIDENTAL O PÉRDIDAS ORGÁNICAS", "ACCIDENTAL", 113300, "DOLARES", "no disponible"),
    ("92834", "MARTHA LETICIA GOSSEREZ ROJAS", "BIENESTAR PRUDENTIAL 10 AÑOS SIN INVALIDEZ (UDIS)", "BASICA", 250000, "UDIS", "PAGO DE PRIMAS EN CURSO"),
    ("92834", "MARTHA LETICIA GOSSEREZ ROJAS", "ENFERMEDADES GRAVES PLUS SIN INVALIDEZ (UDIS)", "SALUD", 27500, "UDIS", "PAGO DE PRIMAS EN CURSO"),
    ("92834", "MARTHA LETICIA GOSSEREZ ROJAS", "INVALIDEZ TOTAL Y PERMANENTE", "INVALIDEZ", 250000, "UDIS", "no disponible"),
    ("92834", "MARTHA LETICIA GOSSEREZ ROJAS", "MUERTE ACCIDENTAL O PÉRDIDAS ORGÁNICAS", "ACCIDENTAL", 250000, "UDIS", "no disponible"),
    ("100631", "ALDO CARDONA GOSSEREZ", "PERSONALIZA PRU 20 AÑOS (UDIS)", "BASICA", 60000, "UDIS", "PAGO DE PRIMAS EN CURSO"),
    ("100631", "ALDO CARDONA GOSSEREZ", "EXCESO PERSONALIZA PRU 20 AÑOS (UDIS)", "EXCESO", 0, "UDIS", "PAGO DE PRIMAS EN CURSO"),
    ("100631", "ALDO CARDONA GOSSEREZ", "MUERTE ACCIDENTAL O PÉRDIDAS ORGÁNICAS", "ACCIDENTAL", 370000, "UDIS", "no disponible"),
]

wb = Workbook()


def style_header(ws, row, ncols, width_map):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=F, bold=True, color=WHITE, size=9)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 30
    for col, w in width_map.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


# ------------------------------------------------------------------ RESUMEN
r = wb.active
r.title = "Resumen"
r.sheet_view.showGridLines = False
r.column_dimensions["A"].width = 46
r.column_dimensions["B"].width = 22
r.column_dimensions["C"].width = 40

r["A1"] = "Inventario de pólizas y sumas aseguradas"
r["A1"].font = Font(name=F, bold=True, size=15, color=INK)
r["A2"] = "Daniel Cardona Sánchez  ·  Prudential Seguros México, S.A. de C.V."
r["A2"].font = Font(name=F, size=10, color="FF6B7280")
r["A3"] = "Fecha de corte: 28 de julio de 2026"
r["A3"].font = Font(name=F, size=10, color="FF6B7280")


def section(cell, text):
    r[cell] = text
    r[cell].font = Font(name=F, bold=True, size=10, color=INK)
    r[cell].fill = SEC_FILL
    row = int(cell[1:])
    for c in "ABC":
        r[f"{c}{row}"].fill = SEC_FILL


section("A5", "PARÁMETROS DE VALUACIÓN  (editar estas celdas)")
labels = [
    ("A6", "Valor de la UDI en pesos", "B6", 8.830830, RATE,
     "Banxico. Actualízalo el día que uses el archivo."),
    ("A7", "Tipo de cambio USD/MXN (FIX)", "B7", 17.5130, RATE,
     "DOF, publicado para el 24/07/2026."),
]
for la, lt, cb, val, fmt, note in labels:
    r[la] = lt
    r[la].font = Font(name=F, size=10)
    r[cb] = val
    r[cb].font = Font(name=F, size=10, bold=True, color=BLUE)
    r[cb].fill = INPUT_FILL
    r[cb].number_format = fmt
    r[cb].border = BOX
    nc = "C" + cb[1:]
    r[nc] = note
    r[nc].font = Font(name=F, size=9, italic=True, color="FF6B7280")

r["A8"] = "Todas las cifras en pesos del archivo se recalculan a partir de estas dos celdas."
r["A8"].font = Font(name=F, size=9, italic=True, color="FF6B7280")

section("A10", "NÚMERO DE PÓLIZAS")
conteo = [
    ("A11", "Pólizas del titular (Daniel Cardona Sánchez)", "B11",
     '=COUNTIF(Polizas!$D$2:$D$7,"Titular")'),
    ("A12", "Pólizas del grupo familiar", "B12",
     '=COUNTIF(Polizas!$D$2:$D$7,"Grupo familiar")'),
    ("A13", "Total de pólizas vigentes", "B13", "=B11+B12"),
]
for la, lt, cb, f in conteo:
    r[la] = lt
    r[la].font = Font(name=F, size=10, bold=(la == "A13"))
    r[cb] = f
    r[cb].font = Font(name=F, size=10, bold=(cb == "B13"))
    r[cb].number_format = "#,##0"
    r[cb].alignment = Alignment(horizontal="right")
r["A13"].border = TOPRULE
r["B13"].border = TOPRULE

section("A15", "SUMA ASEGURADA DEL TITULAR  (pesos)")
tit = [
    ("A16", "Fallecimiento por cualquier causa", "B16",
     '=SUMIF(Polizas!$D$2:$D$7,"Titular",Polizas!$M$2:$M$7)',
     "Sólo cobertura básica. Es la protección de piso real."),
    ("A17", "Fallecimiento accidental", "B17",
     '=SUMIF(Polizas!$D$2:$D$7,"Titular",Polizas!$P$2:$P$7)',
     "Básica más muerte accidental. Aplica sólo si el fallecimiento es accidental."),
    ("A18", "Invalidez total y permanente", "B18",
     '=SUMIF(Polizas!$D$2:$D$7,"Titular",Polizas!$N$2:$N$7)',
     "Beneficio independiente. No se suma al fallecimiento."),
]
section("A20", "SUMA ASEGURADA DEL GRUPO FAMILIAR  (pesos)")
fam = [
    ("A21", "Fallecimiento por cualquier causa", "B21",
     '=SUMIF(Polizas!$D$2:$D$7,"Grupo familiar",Polizas!$M$2:$M$7)', ""),
    ("A22", "Fallecimiento accidental", "B22",
     '=SUMIF(Polizas!$D$2:$D$7,"Grupo familiar",Polizas!$P$2:$P$7)', ""),
    ("A23", "Invalidez total y permanente", "B23",
     '=SUMIF(Polizas!$D$2:$D$7,"Grupo familiar",Polizas!$N$2:$N$7)', ""),
]
section("A25", "TOTAL CONSOLIDADO  (pesos)")
tot = [
    ("A26", "Fallecimiento por cualquier causa", "B26", "=B16+B21", ""),
    ("A27", "Fallecimiento accidental", "B27", "=B17+B22", ""),
]
for group in (tit, fam, tot):
    for la, lt, cb, f, note in group:
        r[la] = lt
        r[la].font = Font(name=F, size=10)
        r[cb] = f
        r[cb].font = Font(name=F, size=11, bold=True)
        r[cb].number_format = MXN
        r[cb].border = BOX
        if note:
            nc = "C" + cb[1:]
            r[nc] = note
            r[nc].font = Font(name=F, size=9, italic=True, color="FF6B7280")
for cb in ("B26", "B27"):
    r[cb].fill = PatternFill("solid", fgColor="FFFEF7EC")

r["A29"] = "Notas"
r["A29"].font = Font(name=F, bold=True, size=10)
notas = [
    "Las sumas aseguradas no se acumulan de forma plana. La cobertura por muerte accidental "
    "se suma a la básica únicamente cuando el fallecimiento es de origen accidental.",
    "Las coberturas de invalidez y de salud son beneficios independientes y no se suman al "
    "beneficio por fallecimiento.",
    "Fuente de los datos: Portal de Agentes Prudential, ruta Consultas > Pólizas, pestaña "
    "Coberturas de la carátula, consultada el 28/07/2026.",
    "Pendientes: designación de beneficiarios, valores de rescate y confirmar si el plazo de "
    "5 años de la póliza 102264 es periodo de pago o vigencia de la cobertura.",
]
for i, n in enumerate(notas):
    c = r.cell(row=30 + i, column=1, value="· " + n)
    c.font = Font(name=F, size=9, color="FF6B7280")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    r.merge_cells(start_row=30 + i, start_column=1, end_row=30 + i, end_column=3)
    r.row_dimensions[30 + i].height = 26

# ------------------------------------------------------------------ POLIZAS
p = wb.create_sheet("Polizas")
p.sheet_view.showGridLines = False
cols = ["Póliza", "Contratante", "Asegurado", "Titularidad", "Producto", "Emisión",
        "Estatus", "Moneda", "SA básica", "SA invalidez", "SA accidental", "SA salud",
        "Básica MXN", "Invalidez MXN", "Accidental MXN", "Total accidental MXN"]
for i, h in enumerate(cols, start=1):
    p.cell(row=1, column=i, value=h)
widths = {"A": 9, "B": 30, "C": 30, "D": 14, "E": 40, "F": 11, "G": 11, "H": 10,
          "I": 13, "J": 13, "K": 14, "L": 11, "M": 14, "N": 14, "O": 15, "P": 19}
style_header(p, 1, len(cols), widths)

for i, row in enumerate(POLIZAS, start=2):
    for j, v in enumerate(row, start=1):
        c = p.cell(row=i, column=j, value=v)
        c.font = Font(name=F, size=9)
        c.border = BOX
        if i % 2 == 0:
            c.fill = BAND
        if j >= 9:
            c.number_format = NUM
            c.alignment = Alignment(horizontal="right")
    fx = f'=IF($H{i}="UDIS",{{col}}{i}*Resumen!$B$6,{{col}}{i}*Resumen!$B$7)'
    p[f"M{i}"] = fx.format(col="I")
    p[f"N{i}"] = fx.format(col="J")
    p[f"O{i}"] = fx.format(col="K")
    p[f"P{i}"] = f"=M{i}+O{i}"
    for col in "MNOP":
        c = p[f"{col}{i}"]
        c.font = Font(name=F, size=9, bold=(col == "P"))
        c.number_format = MXN
        c.border = BOX
        if i % 2 == 0:
            c.fill = BAND

tr = len(POLIZAS) + 2
p.cell(row=tr, column=1, value="TOTAL").font = Font(name=F, bold=True, size=9)
for col in "MNOP":
    c = p[f"{col}{tr}"]
    c.value = f"=SUM({col}2:{col}{len(POLIZAS)+1})"
    c.font = Font(name=F, bold=True, size=9)
    c.number_format = MXN
    c.border = TOPRULE
p[f"A{tr}"].border = TOPRULE

p.cell(row=tr + 2, column=1,
       value="Las columnas en pesos se calculan con la UDI y el tipo de cambio de la hoja "
             "Resumen, celdas B6 y B7.").font = Font(name=F, size=9, italic=True,
                                                     color="FF6B7280")
p.auto_filter.ref = f"A1:P{len(POLIZAS)+1}"

# --------------------------------------------------------------- COBERTURAS
d = wb.create_sheet("Coberturas")
d.sheet_view.showGridLines = False
dcols = ["Póliza", "Asegurado", "Cobertura", "Tipo", "Suma asegurada", "Moneda",
         "Estatus cobertura", "Equivalente MXN"]
for i, h in enumerate(dcols, start=1):
    d.cell(row=1, column=i, value=h)
style_header(d, 1, len(dcols),
             {"A": 9, "B": 30, "C": 48, "D": 12, "E": 15, "F": 10, "G": 24, "H": 15})

for i, row in enumerate(COBERTURAS, start=2):
    for j, v in enumerate(row, start=1):
        c = d.cell(row=i, column=j, value=v)
        c.font = Font(name=F, size=9)
        c.border = BOX
        if i % 2 == 0:
            c.fill = BAND
        if j == 5:
            c.number_format = NUM
            c.alignment = Alignment(horizontal="right")
    c = d[f"H{i}"]
    c.value = f'=IF($F{i}="UDIS",E{i}*Resumen!$B$6,E{i}*Resumen!$B$7)'
    c.font = Font(name=F, size=9)
    c.number_format = MXN
    c.border = BOX
    if i % 2 == 0:
        c.fill = BAND

d.auto_filter.ref = f"A1:H{len(COBERTURAS)+1}"
d.cell(row=len(COBERTURAS) + 3, column=1,
       value="Detalle completo de las 21 coberturas. Filtra por la columna Tipo para aislar "
             "BASICA, INVALIDEZ, ACCIDENTAL, SALUD o EXCESO.").font = Font(
    name=F, size=9, italic=True, color="FF6B7280")

wb.save(OUT)
print("OK ->", OUT)
